import os
import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from .db import SessionLocal
from .models import Invoice, Concept, Job
HEADERS = [
"NUM_CTA_BCO",
"NUM_OPERACION",
"FECHA_OPE_BANCO",
"VALOR_OPE_BANCO",
"tc",
"f_id",
"f_rfc",
"f_razon_social",
"f_folio",
"f_cert_cfdi",
"f_importe",
"f_descuento",
"f_ieps",
"f_iva",
"f_ret_iva",
"f_ret_isr",
"f_total",
"f_mn_cfdi",
"f_mpago_cfdi",
"f_fpago_cfdi",
"f_uso_cfdi",
"f_uuid",

"p_uuid",
"p_id",
"p_fecha_pago",
"p_forma_pago",
"p_moneda",
"p_num_parcialidad",
"p_imp_saldo_ant",
"p_imp_pagado",
"p_sal_insoluto",
"p_fecha_emision",
"pagos_agrupados",

"d_cant",
"d_desc",
"d_pu",
"d_importe",
"d_dcto",
"d_ieps",
"d_iva",
"d_ret_iva",
"d_ret_isr",
"d_total",
"d_t_ieps",
"d_t_iva",
"d_t_ret_iva",
"d_t_ret_isr",
"d_sat_prod",
"d_sat_unid"
]
def row(i,c,aux_map):

    uuid = str(i.uuid or "").strip().upper()

    aux = aux_map.get(uuid, {})

    return [
        aux.get("NUM_CTA_BCO",""),
        aux.get("NUM_OPERACION",""),
        aux.get("FECHA_OPE_BANCO",""),
        aux.get("VALOR_OPE_BANCO",""),
        aux.get("tc",""),
        i.id,
        i.issuer_rfc,
        i.issuer_name,
        i.folio,
        i.issue_date,
        i.subtotal,
        i.discount,
        c.ieps_transferred,
        c.vat_transferred,
        c.vat_withheld,
        c.isr_withheld,
        i.total,
        i.currency,
        i.payment_method,
        i.payment_form,
        i.cfdi_use,
        i.uuid,

        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",

        c.quantity,
        c.description,
        c.unit_value,
        c.amount,
        c.discount,
        c.ieps_transferred,
        c.vat_transferred,
        c.vat_withheld,
        c.isr_withheld,

        (
            (c.amount or 0)
            - (c.discount or 0)
            + (c.vat_transferred or 0)
            + (c.ieps_transferred or 0)
            - (c.vat_withheld or 0)
            - (c.isr_withheld or 0)
        ),

        c.ieps_transferred,
        c.vat_transferred,
        c.vat_withheld,
        c.isr_withheld,

        c.product_key,
        c.unit_key

    ]

def _find_uuid_column(columns):
    """Busca una columna que sea 'UUID' sin importar mayusculas/espacios."""
    for col in columns:
        if str(col).strip().upper() == "UUID":
            return col
    return None

def _load_aux_map(job, aux_map, warnings):
    if not job.aux_file:
        return

    try:
        df = pd.read_excel(job.aux_file)
    except Exception as e:
        warnings.append(f"Job {job.id}: no se pudo leer el Excel auxiliar ({e})")
        return

    uuid_col = _find_uuid_column(df.columns)

    if uuid_col is None:
        warnings.append(
            f"Job {job.id}: el Excel auxiliar no tiene una columna 'UUID'. "
            f"Columnas encontradas: {list(df.columns)}"
        )
        return

    try:
        df[uuid_col] = df[uuid_col].astype(str).str.upper().str.strip()
        aux_map.update(df.set_index(uuid_col).to_dict("index"))
    except Exception as e:
        warnings.append(f"Job {job.id}: error al procesar el Excel auxiliar ({e})")


def create_export(job_id: int | None = None):

    db = SessionLocal()

    jobs = db.query(Job).all() if job_id is None else [db.get(Job, job_id)]

    aux_map = {}
    warnings = []

    for job in jobs:
        if not job:
            continue
        _load_aux_map(job, aux_map, warnings)

    if warnings:
        # Se imprime en los logs de Render para poder diagnosticar sin adivinar
        print("AVISOS create_export:\n" + "\n".join(warnings))

    out = Path(os.getenv("DATA_DIR", "/tmp")) / "exports"
    out.mkdir(parents=True, exist_ok=True)

    filename = f"detalle_cfdi_job{job_id}.xlsx" if job_id else "detalle_cfdi_todos.xlsx"
    path = out / filename

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Detalle_CFDI")
    ws.append(HEADERS)

    q = (
        select(Invoice, Concept)
        .join(Concept, Concept.invoice_id == Invoice.id)
        .order_by(Invoice.id, Concept.line_no)
    )

    if job_id is not None:
        q = q.where(Invoice.job_id == job_id)

    for i, c in db.execute(q).yield_per(2000):
        ws.append(row(i, c, aux_map))

    db.close()

    wb.save(path)

    return path
